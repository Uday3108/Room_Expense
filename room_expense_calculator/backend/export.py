"""
Excel export using openpyxl.
Produces a workbook with four sheets: Daily, Fixed, Shopping, Summary.
"""

import io
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ── Helpers ──────────────────────────────────────────────────────────────────

HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
HEADER_FONT = Font(color="FFFFFF", bold=True)
BORDER_SIDE = Side(style="thin", color="E5E7EB")
CELL_BORDER = Border(
    left=BORDER_SIDE, right=BORDER_SIDE,
    top=BORDER_SIDE, bottom=BORDER_SIDE,
)


def _style_header_row(ws, row: int, num_cols: int):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = CELL_BORDER


def _style_data_cell(cell, bold=False):
    cell.border = CELL_BORDER
    cell.alignment = Alignment(vertical="center")
    if bold:
        cell.font = Font(bold=True)


def _auto_width(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)


# ── Sheet builders ───────────────────────────────────────────────────────────

def _build_daily_sheet(ws, expenses):
    headers = ["ID", "Date", "Amount (₹)", "Paid By", "Category", "Description"]
    ws.append(headers)
    _style_header_row(ws, 1, len(headers))
    for exp in expenses:
        row = [exp.id, str(exp.date), exp.amount, exp.paid_by,
               exp.category, exp.description or ""]
        ws.append(row)
        for col in range(1, len(headers) + 1):
            _style_data_cell(ws.cell(row=ws.max_row, column=col))
    _auto_width(ws)


def _build_fixed_sheet(ws, expenses):
    headers = ["ID", "Date", "Amount (₹)", "Paid By", "Expense Type"]
    ws.append(headers)
    _style_header_row(ws, 1, len(headers))
    for exp in expenses:
        row = [exp.id, str(exp.date), exp.amount, exp.paid_by, exp.expense_type]
        ws.append(row)
        for col in range(1, len(headers) + 1):
            _style_data_cell(ws.cell(row=ws.max_row, column=col))
    _auto_width(ws)


def _build_shopping_sheet(ws, expenses):
    headers = ["ID", "Date", "Amount (₹)", "Paid By", "Store", "Description"]
    ws.append(headers)
    _style_header_row(ws, 1, len(headers))
    for exp in expenses:
        row = [exp.id, str(exp.date), exp.amount, exp.paid_by,
               exp.store_name, exp.description or ""]
        ws.append(row)
        for col in range(1, len(headers) + 1):
            _style_data_cell(ws.cell(row=ws.max_row, column=col))
    _auto_width(ws)


def _build_summary_sheet(ws, dashboard):
    ws.append(["Summary"])
    ws["A1"].font = Font(bold=True, size=14, color="1E3A5F")
    ws.append([])
    ws.append(["Category", "Amount (₹)"])
    _style_header_row(ws, 3, 2)
    rows = [
        ("Total Daily Expenses", dashboard["total_daily"]),
        ("Total Fixed Expenses", dashboard["total_fixed"]),
        ("Total Shopping Expenses", dashboard["total_shopping"]),
        ("Grand Total", dashboard["grand_total"]),
        ("Per Person Share", dashboard["per_person_share"]),
    ]
    for label, val in rows:
        ws.append([label, val])
        for col in range(1, 3):
            _style_data_cell(ws.cell(row=ws.max_row, column=col))

    ws.append([])
    ws.append(["Member Balances"])
    ws[f"A{ws.max_row}"].font = Font(bold=True, size=12, color="1E3A5F")
    ws.append(["Name", "Paid (₹)", "Share (₹)", "Balance (₹)", "Status"])
    _style_header_row(ws, ws.max_row, 5)
    for mb in dashboard["member_balances"]:
        status = f"Receives ₹{mb['balance']:.2f}" if mb["balance"] >= 0 else f"Owes ₹{abs(mb['balance']):.2f}"
        ws.append([mb["name"], mb["paid"], mb["share"], mb["balance"], status])
        for col in range(1, 6):
            _style_data_cell(ws.cell(row=ws.max_row, column=col))

    ws.append([])
    ws.append(["Settlement Transactions"])
    ws[f"A{ws.max_row}"].font = Font(bold=True, size=12, color="1E3A5F")
    ws.append(["From", "To", "Amount (₹)"])
    _style_header_row(ws, ws.max_row, 3)
    for txn in dashboard["settlement"]:
        ws.append([txn["from_member"], txn["to_member"], txn["amount"]])
        for col in range(1, 4):
            _style_data_cell(ws.cell(row=ws.max_row, column=col))
    _auto_width(ws)


# ── Public API ───────────────────────────────────────────────────────────────

def generate_excel(daily, fixed, shopping, dashboard) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)  # remove default blank sheet

    _build_daily_sheet(wb.create_sheet("Daily Expenses"), daily)
    _build_fixed_sheet(wb.create_sheet("Fixed Expenses"), fixed)
    _build_shopping_sheet(wb.create_sheet("Shopping Expenses"), shopping)
    _build_summary_sheet(wb.create_sheet("Summary"), dashboard)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
